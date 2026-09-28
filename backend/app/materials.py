"""Turns an uploaded file (PDF, Word, text, image) into short parts the AI agent can read one at a time.

Nothing here calls the AI: we only extract the text and cut it along headings and paragraphs,
so a 300-page book becomes a list of parts of a few thousand characters each.
Pictures are pulled out too. Where a picture stood, the text gets a marker "[Рисунок N]", so the
agent later knows which part a picture belongs to (see app/vision.py for reading the pictures).
"""

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from app.models import SourceKind

# A part aims for about TARGET characters (roughly one short lesson of source text)
# and never goes over MAX_CHARS. A heading starts a new part once the current one has MIN_CHARS.
TARGET_CHARS = 3000
MAX_CHARS = 6000
MIN_CHARS = 600

EXTENSIONS = {
    ".pdf": SourceKind.pdf,
    ".docx": SourceKind.docx,
    ".txt": SourceKind.text,
    ".md": SourceKind.text,
    ".markdown": SourceKind.text,
    ".png": SourceKind.image,
    ".jpg": SourceKind.image,
    ".jpeg": SourceKind.image,
    ".webp": SourceKind.image,
}

# Pictures smaller than this are icons, bullets or lines, not figures worth explaining.
MIN_IMAGE_BYTES = 3_000
MIN_IMAGE_SIDE = 64
MAX_IMAGES = 200
FIGURE = "[Рисунок {n}]"
FIGURE_RE = re.compile(r"\[Рисунок (\d+)\]")


class MaterialError(ValueError):
    """The file cannot be used; the message is shown to the teacher as is."""


@dataclass
class Paragraph:
    text: str
    is_heading: bool = False


@dataclass
class Chunk:
    heading: str
    text: str


@dataclass
class Image:
    number: int  # the N in "[Рисунок N]", from 1
    data: bytes
    ext: str  # "png", "jpeg", ...


@dataclass
class ParsedMaterial:
    kind: SourceKind
    chunks: list[Chunk]
    images: list[Image]


class _Figures:
    """Collects pictures while the text is read and hands out their markers."""

    def __init__(self) -> None:
        self.images: list[Image] = []

    def add(self, data: bytes, ext: str) -> Paragraph | None:
        if len(self.images) >= MAX_IMAGES or not _big_enough(data):
            return None
        image = Image(number=len(self.images) + 1, data=data, ext=ext.lower().lstrip(".").replace("jpg", "jpeg"))
        self.images.append(image)
        return Paragraph(FIGURE.format(n=image.number))


def _big_enough(data: bytes) -> bool:
    if len(data) < MIN_IMAGE_BYTES:
        return False
    try:
        from PIL import Image as PILImage

        with PILImage.open(io.BytesIO(data)) as img:
            return min(img.size) >= MIN_IMAGE_SIDE
    except Exception:  # Pillow missing or a format it cannot read: keep the picture
        return True


def detect_kind(filename: str) -> SourceKind:
    ext = PurePath(filename).suffix.lower()
    if ext == ".doc":
        raise MaterialError("Старый формат .doc не поддерживается. Сохраните файл в Word как .docx.")
    if ext not in EXTENSIONS:
        raise MaterialError("Поддерживаются PDF, Word (.docx), текст (.txt, .md) и картинки (.png, .jpg, .webp).")
    return EXTENSIONS[ext]


# --- Extraction -----------------------------------------------------------------------

_NUMBERED_HEADING = re.compile(r"^(\d+(\.\d+)*[.)]?|глава|chapter|раздел|section|bob|mavzu)\s+\S", re.IGNORECASE)


def _looks_like_heading(line: str) -> bool:
    """PDF has no heading styles, so short title-like lines are treated as headings."""
    line = line.strip()
    if not line or len(line) > 90 or line[-1] in ".,;:!?":
        return False
    if _NUMBERED_HEADING.match(line):
        return True
    letters = [c for c in line if c.isalpha()]
    return len(letters) >= 4 and all(c.isupper() for c in letters)


def _split_plain(text: str, markdown: bool) -> list[Paragraph]:
    """Blank lines separate paragraphs; '#' lines (Markdown) or title-like lines are headings."""
    paragraphs: list[Paragraph] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            paragraphs.append(Paragraph(" ".join(buffer)))
            buffer.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
        elif markdown and line.startswith("#"):
            flush()
            paragraphs.append(Paragraph(line.lstrip("#").strip(), is_heading=True))
        elif not markdown and not buffer and _looks_like_heading(line):
            paragraphs.append(Paragraph(line, is_heading=True))
        else:
            buffer.append(line)
    flush()
    return [p for p in paragraphs if p.text]


def _decode(data: bytes) -> str:
    # Russian and Uzbek texts from Windows often come in cp1251 rather than UTF-8.
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _extract_pdf(data: bytes, figures: _Figures) -> list[Paragraph]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise MaterialError("PDF защищён паролем. Снимите защиту и загрузите снова.")
        paragraphs: list[Paragraph] = []
        for page in reader.pages:
            paragraphs += _split_plain(page.extract_text() or "", markdown=False)
            # PDF does not say where on the page a picture sits, so its marker goes after the page text.
            for picture in _page_images(page):
                marker = figures.add(picture.data, PurePath(picture.name).suffix or ".png")
                if marker:
                    paragraphs.append(marker)
    except PdfReadError as exc:
        raise MaterialError("Не удалось прочитать PDF: файл повреждён.") from exc
    return paragraphs


def _page_images(page) -> list:
    try:
        return list(page.images)
    except Exception:  # an unusual image encoding must not lose the page text
        return []


def _extract_docx(data: bytes, figures: _Figures) -> list[Paragraph]:
    from docx import Document
    from docx.oxml.ns import qn

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for broken files
        raise MaterialError("Не удалось прочитать файл Word: он повреждён или это не .docx.") from exc

    paragraphs: list[Paragraph] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            style = (p.style.name if p.style is not None else "").lower()
            is_heading = style.startswith(("heading", "title", "заголовок", "название"))
            paragraphs.append(Paragraph(text, is_heading=is_heading))
        for blip in p._p.xpath(".//a:blip"):
            part = doc.part.related_parts.get(blip.get(qn("r:embed")))
            if part is not None:
                marker = figures.add(part.blob, PurePath(part.partname).suffix)
                if marker:
                    paragraphs.append(marker)
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            paragraphs.append(Paragraph("\n".join(rows)))
    return paragraphs


def extract_paragraphs(kind: SourceKind, filename: str, data: bytes, figures: _Figures | None = None) -> list[Paragraph]:
    figures = figures if figures is not None else _Figures()
    if kind == SourceKind.pdf:
        paragraphs = _extract_pdf(data, figures)
    elif kind == SourceKind.docx:
        paragraphs = _extract_docx(data, figures)
    elif kind == SourceKind.image:
        figures.images.append(Image(number=1, data=data, ext=PurePath(filename).suffix.lower().lstrip(".").replace("jpg", "jpeg")))
        paragraphs = [Paragraph(FIGURE.format(n=1))]
    else:
        markdown = PurePath(filename).suffix.lower() in (".md", ".markdown")
        paragraphs = _split_plain(_decode(data), markdown=markdown)
    if not paragraphs:
        raise MaterialError("В файле нет ни текста, ни рисунков.")
    return paragraphs


# --- Splitting into parts -------------------------------------------------------------

_SENTENCE_END = re.compile(r"(?<=[.!?…])\s+")


def _split_long(text: str) -> list[str]:
    """Cuts one oversized paragraph at sentence ends, or hard at MAX_CHARS as a last resort."""
    pieces: list[str] = []
    current = ""
    for sentence in _SENTENCE_END.split(text):
        while len(sentence) > MAX_CHARS:
            pieces.append(sentence[:MAX_CHARS])
            sentence = sentence[MAX_CHARS:]
        if current and len(current) + 1 + len(sentence) > TARGET_CHARS:
            pieces.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return pieces


def split_into_chunks(paragraphs: list[Paragraph]) -> list[Chunk]:
    chunks: list[Chunk] = []
    heading = ""
    body: list[str] = []
    size = 0

    def flush() -> None:
        nonlocal body, size
        if body:
            chunks.append(Chunk(heading=heading, text="\n\n".join(body)))
        body, size = [], 0

    for p in paragraphs:
        if p.is_heading:
            if size >= MIN_CHARS:
                flush()
            if body:
                # A short section stays with the previous one; keep its heading visible in the text.
                body.append(p.text)
                size += len(p.text)
            else:
                heading = p.text
            continue
        for piece in _split_long(p.text) if len(p.text) > MAX_CHARS else [p.text]:
            if body and size + len(piece) > TARGET_CHARS:
                flush()
            body.append(piece)
            size += len(piece)
    flush()

    for chunk in chunks:
        if not chunk.heading:
            chunk.heading = chunk.text[:60].split("\n")[0]
        chunk.heading = chunk.heading[:300]
    return chunks


def read_material(filename: str, data: bytes) -> ParsedMaterial:
    kind = detect_kind(filename)
    figures = _Figures()
    chunks = split_into_chunks(extract_paragraphs(kind, filename, data, figures))
    return ParsedMaterial(kind=kind, chunks=chunks, images=figures.images)


def parse_material(filename: str, data: bytes) -> tuple[SourceKind, list[Chunk]]:
    parsed = read_material(filename, data)
    return parsed.kind, parsed.chunks


def figure_numbers(text: str) -> list[int]:
    return [int(n) for n in FIGURE_RE.findall(text)]
