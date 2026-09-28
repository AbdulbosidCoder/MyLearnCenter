import io
import os
import shutil

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app import agent, rag, vision
from app.config import get_settings
from app.db import SessionLocal
from app.materials import read_material
from tests.conftest import ADMIN_ID, auth, indexed


def picture(size: int = 160, fmt: str = "PNG") -> bytes:
    """Noise, so the file is big enough to count as a figure and not an icon."""
    image = Image.frombytes("RGB", (size, size), os.urandom(size * size * 3))
    out = io.BytesIO()
    image.save(out, fmt)
    return out.getvalue()


def word_with_picture() -> bytes:
    doc = Document()
    doc.add_heading("Mean", level=1)
    doc.add_paragraph("The mean is the sum divided by the count. " * 20)
    doc.add_picture(io.BytesIO(picture()))
    doc.add_paragraph("Outliers move the mean a lot. " * 20)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_pictures_come_out_of_word_and_pdf_with_markers():
    parsed = read_material("stats.docx", word_with_picture())
    assert [(i.number, i.ext) for i in parsed.images] == [(1, "png")]
    assert "[Рисунок 1]" in parsed.chunks[0].text
    assert parsed.chunks[0].text.index("[Рисунок 1]") < parsed.chunks[0].text.index("Outliers")

    # A scanned PDF: pages are pictures without text.
    out = io.BytesIO()
    Image.open(io.BytesIO(picture(300))).save(out, "PDF")
    scan = read_material("scan.pdf", out.getvalue())
    assert len(scan.images) == 1 and scan.chunks[0].text == "[Рисунок 1]"

    # Icons are skipped.
    assert read_material("icon.docx", word_with_icon()).images == []


def word_with_icon() -> bytes:
    doc = Document()
    doc.add_paragraph("Text")
    doc.add_picture(io.BytesIO(picture(size=20)))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def upload(client, filename: str, data: bytes, module_id: int | None = None) -> dict:
    extra = {"module_id": str(module_id)} if module_id else {}
    r = await client.post("/api/materials", files={"file": (filename, data)}, data=extra, headers=auth(ADMIN_ID))
    assert r.status_code == 201, r.text
    assert r.json()["index_status"] == "pending"
    await indexed()
    return (await client.get(f"/api/materials/{r.json()['id']}", headers=auth(ADMIN_ID))).json()


async def test_pictures_are_read_saved_and_searchable(client):
    material = await upload(client, "stats.docx", word_with_picture())
    assert material["index_status"] == "ready" and material["images_done"] == 1 and material["image_count"] == 1
    [image] = material["images"]
    assert image["caption"] == "a bar chart of mean values" and image["ocr_text"] == "Mean = 5"
    assert (await client.get(image["url"])).status_code == 200

    async with SessionLocal() as session:
        found = await rag.search(session, "bar chart", document_ids=[material["id"]])
    assert found[0][0].image_id == image["id"]

    # Deleting the material removes its pictures.
    assert (await client.delete(f"/api/materials/{material['id']}", headers=auth(ADMIN_ID))).status_code == 204
    assert not (rag.media_root() / "materials" / str(material["id"])).exists()


async def test_scan_is_split_by_what_ocr_read(client, monkeypatch):
    page_text = "Dispersiya - bu ma'lumotlarning tarqalishi.\n\n" + "Standart og'ish dispersiyadan ildiz. " * 40
    monkeypatch.setattr(vision, "read_image", lambda data: vision.ImageText("a page of text", page_text))
    out = io.BytesIO()
    Image.open(io.BytesIO(picture(300))).save(out, "PDF")
    material = await upload(client, "scan.pdf", out.getvalue())
    assert "Dispersiya" in material["chunks"][0]["text"]
    assert material["chunks"][-1]["text"].endswith("[Рисунок 1]")


async def test_agent_puts_pictures_into_lessons(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "anthropic_api_key", "test-key")
    prompts = []

    async def ask(prompt, schema):
        prompts.append(prompt)
        if schema is agent.PLAN_SCHEMA:
            return {"title": "Mean", "description": "", "order": [1]}
        card = {"heading": "Idea", "text": "See the picture."}
        return {"lessons": [{"title": "Mean", "level": 1, "cards": [card, card], "figures": [1, 7], "questions": []}]}

    monkeypatch.setattr(agent, "ask_claude", ask)
    material = await upload(client, "stats.docx", word_with_picture())
    await client.post(f"/api/materials/{material['id']}/generate", headers=auth(ADMIN_ID))
    await __import__("asyncio").gather(*agent._running)

    assert "[Рисунок 1]\nНа картинке: a bar chart of mean values\nТекст на картинке: Mean = 5" in prompts[0]
    done = (await client.get(f"/api/materials/{material['id']}", headers=auth(ADMIN_ID))).json()
    module = (await client.get(f"/api/modules/{done['module_id']}", headers=auth(ADMIN_ID))).json()
    lesson = (await client.get(f"/api/lessons/{module['lessons'][0]['id']}", headers=auth(ADMIN_ID))).json()
    # Picture 7 does not exist and is dropped.
    assert [(b["type"], b["content"]) for b in lesson["blocks"]][1] == ("image", material["images"][0]["url"])
    assert len(lesson["blocks"]) == 3


async def test_helper_answers_with_a_picture(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "anthropic_api_key", "test-key")
    asked = []

    async def ask_text(system, prompt):
        asked.append(prompt)
        return "The mean is **5**."

    monkeypatch.setattr(agent, "ask_text", ask_text)
    modules = (await client.get("/api/modules", headers=auth(ADMIN_ID))).json()
    stats = next(m for m in modules if m["title"] == "Статистика")
    await upload(client, "stats.docx", word_with_picture(), module_id=stats["id"])
    lesson_id = (await client.get(f"/api/modules/{stats['id']}", headers=auth(ADMIN_ID))).json()["lessons"][0]["id"]

    r = await client.post(
        "/api/ai/ask",
        data={"question": "Что показывает график?", "lesson_id": str(lesson_id)},
        files={"image": ("board.png", picture(), "image/png")},
        headers=auth(500),
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answer"] == "The mean is **5**."
    assert "a bar chart of mean values" in body["image_text"]
    assert any(s["image_url"] for s in body["sources"])
    assert "Среднее и медиана" in asked[0] and "Что показывает график?" in asked[0]

    bad = await client.post("/api/ai/ask", data={"question": ""}, headers=auth(500))
    assert bad.status_code == 422
    not_image = await client.post("/api/ai/ask", files={"image": ("a.txt", b"x", "text/plain")}, headers=auth(500))
    assert not_image.status_code == 422


async def test_students_do_not_search_unpublished_materials(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "anthropic_api_key", "test-key")

    async def ask_text(system, prompt):
        return "ok"

    monkeypatch.setattr(agent, "ask_text", ask_text)
    await upload(client, "stats.docx", word_with_picture())  # no theme yet: a draft for teachers
    student = (await client.post("/api/ai/ask", data={"question": "bar chart mean"}, headers=auth(500))).json()
    teacher = (await client.post("/api/ai/ask", data={"question": "bar chart mean"}, headers=auth(ADMIN_ID))).json()
    assert student["sources"] == [] and teacher["sources"]


async def test_ai_status(client):
    status = (await client.get("/api/ai/status", headers=auth(500))).json()
    assert status["embeddings"] is False and status["claude"] is False


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract is not installed")
def test_tesseract_reads_text_from_a_picture():
    vision.ocr_languages.cache_clear()
    image = Image.new("RGB", (900, 160), "white")
    font = ImageFont.load_default(size=48)
    ImageDraw.Draw(image).text((20, 50), "Mean value 42", fill="black", font=font)
    out = io.BytesIO()
    image.save(out, "PNG")
    assert "Mean value 42" in vision.read_text(out.getvalue())
