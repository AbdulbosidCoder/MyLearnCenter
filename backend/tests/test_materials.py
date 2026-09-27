import io

from docx import Document

from app.materials import MAX_CHARS, TARGET_CHARS, Paragraph, parse_material, split_into_chunks
from tests.conftest import ADMIN_ID, auth


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal text PDF (Helvetica, one line per string), enough for pypdf to extract."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for lines in pages:
        ops = ["BT /F1 12 Tf 50 750 Td 14 TL"] + [f"({line}) Tj T*" for line in lines] + ["ET"]
        stream = "\n".join(ops)
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objects)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            "/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Mean and median", level=1)
    doc.add_paragraph("The mean is the sum divided by the count. " * 30)
    doc.add_heading("Variance", level=1)
    doc.add_paragraph("Variance measures the spread of the data around the mean. " * 30)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_sections_become_parts_and_long_text_is_cut():
    long_section = [Paragraph(f"Paragraph {i}. " + "word " * 150) for i in range(20)]
    paragraphs = [Paragraph("Intro", True), Paragraph("x " * 400), Paragraph("Big topic", True), *long_section]
    chunks = split_into_chunks(paragraphs)

    assert chunks[0].heading == "Intro"
    assert all(c.heading == "Big topic" for c in chunks[1:])
    assert len(chunks) > 3
    assert all(len(c.text) <= TARGET_CHARS + MAX_CHARS for c in chunks)


def test_one_huge_paragraph_is_split_at_sentences():
    chunks = split_into_chunks([Paragraph("A short sentence here. " * 2000)])
    assert len(chunks) > 5
    assert all(len(c.text) <= TARGET_CHARS for c in chunks)
    assert all(c.text.endswith(".") for c in chunks)


def test_markdown_headings_and_short_sections_merge():
    text = "# Intro\n\nTiny.\n\n# Next\n\nAlso tiny.\n".encode()
    kind, chunks = parse_material("notes.md", text)
    assert kind == "text"
    assert len(chunks) == 1 and chunks[0].heading == "Intro"
    assert "Next" in chunks[0].text


def test_cp1251_text_is_decoded():
    _, chunks = parse_material("lesson.txt", "Среднее значение выборки".encode("cp1251"))
    assert chunks[0].text == "Среднее значение выборки"


def test_pdf_pages_are_read():
    pdf = make_pdf([["1. Statistics", "Data has a centre and a spread."], ["2. Regression", "A line fits the points."]])
    kind, chunks = parse_material("book.pdf", pdf)
    text = "\n".join(c.text for c in chunks)
    assert kind == "pdf"
    assert chunks[0].heading == "1. Statistics"
    assert "A line fits the points." in text


async def test_teacher_uploads_word_file(client):
    r = await client.post(
        "/api/materials",
        files={"file": ("stats.docx", make_docx())},
        data={"title": "Статистика"},
        headers=auth(ADMIN_ID),
    )
    assert r.status_code == 201, r.text
    material = r.json()
    assert material["kind"] == "docx" and material["status"] == "parsed"
    assert material["chunk_count"] == 2

    listed = (await client.get("/api/materials", headers=auth(ADMIN_ID))).json()
    assert [m["title"] for m in listed] == ["Статистика"]

    detail = (await client.get(f"/api/materials/{material['id']}", headers=auth(ADMIN_ID))).json()
    assert [c["heading"] for c in detail["chunks"]] == ["Mean and median", "Variance"]

    assert (await client.delete(f"/api/materials/{material['id']}", headers=auth(ADMIN_ID))).status_code == 204
    assert (await client.get(f"/api/materials/{material['id']}", headers=auth(ADMIN_ID))).status_code == 404


async def test_bad_files_are_explained(client):
    h = auth(ADMIN_ID)
    r = await client.post("/api/materials", files={"file": ("old.doc", b"x")}, headers=h)
    assert r.status_code == 422 and ".docx" in r.json()["detail"]
    r = await client.post("/api/materials", files={"file": ("scan.pdf", make_pdf([[]]))}, headers=h)
    assert r.status_code == 422 and "скан" in r.json()["detail"]
    r = await client.post("/api/materials", files={"file": ("broken.docx", b"not a zip")}, headers=h)
    assert r.status_code == 422


async def test_students_cannot_upload_or_read_materials(client):
    r = await client.post("/api/materials", files={"file": ("a.txt", b"hello")}, headers=auth(500))
    assert r.status_code == 403
    assert (await client.get("/api/materials", headers=auth(500))).status_code == 403
